import os
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

import django
import io

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from apps.business.models import Business
from apps.crm.models import Client
from apps.crm.views import parse_and_import_clients, client_template_download_view
from django.test import RequestFactory
from django.contrib.auth import get_user_model

def run_tests():
    print("=== INICIANDO PRUEBAS DE IMPORTACIÓN DE CLIENTES ===")
    
    # 1. Obtener o crear negocio de prueba
    business = Business.objects.first()
    if not business:
        business = Business.objects.create(
            name="Negocio de Prueba Importación",
            slug="test-import-business",
            phone="+5491111111111",
            email="test@business.com"
        )
    print(f"[OK] Negocio para test: {business.name} (ID: {business.id})")
    
    # Limpieza previa de clientes de prueba
    test_phones = ['+541100010001', '+5491100020002', '+5491100030003', '+5491100040004']
    Client.objects.filter(business=business, phone__in=test_phones).delete()
    
    # 2. Test Descarga de Plantilla Excel
    print("\n--- Test 1: Generación y Descarga de Plantilla Excel ---")
    factory = RequestFactory()
    request = factory.get('/crm/clientes/plantilla/')
    request.user = get_user_model().objects.first() or get_user_model().objects.create_user(username='testadmin', password='testpassword')
    request.current_business = business
    
    response = client_template_download_view(request)
    assert response.status_code == 200, f"Error HTTP status {response.status_code}"
    assert 'spreadsheetml' in response['Content-Type'], f"Content type incorrecto: {response['Content-Type']}"
    assert 'Plantilla_Importar_Clientes.xlsx' in response['Content-Disposition']
    
    # Verificar contenido del archivo Excel generado
    template_wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Clientes" in template_wb.sheetnames, "Falta hoja Clientes"
    assert "Instrucciones" in template_wb.sheetnames, "Falta hoja Instrucciones"
    ws = template_wb["Clientes"]
    headers = [cell.value for cell in ws[1]]
    print(f"Encabezados plantilla: {headers}")
    assert "Nombre *" in headers[0]
    print("[PASSED] Test 1: Plantilla generada correctamente con hojas y encabezados válidos.")

    # 3. Test Importación desde archivo Excel (.xlsx)
    print("\n--- Test 2: Importación desde archivo .xlsx ---")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Nombre', 'Apellido', 'WhatsApp', 'Email', 'DNI', 'Direccion', 'Estado', 'Notas'])
    ws.append(['Valentina', 'Ríos', '1100010001', 'valentina.rios@test.com', '40111222', 'Palermo 123', 'VIP', 'Cliente preferencial'])
    ws.append(['Esteban Quito', '', '+5491100020002', 'esteban@test.com', '35999888', 'Av Santa Fe 500', 'Activo', 'Sin apellido explícito'])
    
    excel_io = io.BytesIO()
    wb.save(excel_io)
    excel_io.seek(0)
    
    uploaded_file = SimpleUploadedFile("clientes_test.xlsx", excel_io.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    
    res = parse_and_import_clients(uploaded_file, business, default_prefix='+54', duplicate_mode='skip')
    print(f"Resultado importación .xlsx: {res}")
    assert res['created'] == 2, f"Se esperaban 2 creados, se obtuvieron: {res['created']}"
    
    # Validar cliente 1
    c1 = Client.objects.get(business=business, phone='+541100010001')
    assert c1.first_name == 'Valentina'
    assert c1.last_name == 'Ríos'
    assert c1.status == 'VIP'
    assert c1.tax_id == '40111222'
    print(f"[OK] Cliente 1 creado con éxito: {c1.full_name} ({c1.phone}) - {c1.status}")
    
    # Validar cliente 2 con separación automática de nombre
    c2 = Client.objects.get(business=business, phone='+5491100020002')
    assert c2.first_name == 'Esteban'
    assert c2.last_name == 'Quito'
    print(f"[OK] Cliente 2 separado automáticamente: Nombre='{c2.first_name}', Apellido='{c2.last_name}'")
    print("[PASSED] Test 2: Importación de archivo .xlsx exitosa.")

    # 4. Test Duplicados: Modo Omitir ('skip')
    print("\n--- Test 3: Duplicados con modo 'skip' ---")
    excel_io.seek(0)
    uploaded_dup_file = SimpleUploadedFile("clientes_dup.xlsx", excel_io.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    res_skip = parse_and_import_clients(uploaded_dup_file, business, default_prefix='+54', duplicate_mode='skip')
    print(f"Resultado duplicados en skip: {res_skip}")
    assert res_skip['created'] == 0, "No debió crear registros nuevos"
    assert res_skip['skipped'] == 2, f"Debió omitir 2 duplicados, omitió: {res_skip['skipped']}"
    print("[PASSED] Test 3: Duplicados omitidos correctamente en modo 'skip'.")

    # 5. Test Duplicados: Modo Actualizar ('update')
    print("\n--- Test 4: Duplicados con modo 'update' ---")
    wb_update = openpyxl.Workbook()
    ws_update = wb_update.active
    ws_update.append(['Nombre', 'Apellido', 'WhatsApp', 'Email', 'DNI', 'Direccion', 'Estado', 'Notas'])
    ws_update.append(['Valentina', 'Ríos', '1100010001', 'valentina.rios@test.com', '40111222', 'NUEVA DIRECCION 999', 'VIP', 'Nota adicional'])
    
    excel_update_io = io.BytesIO()
    wb_update.save(excel_update_io)
    excel_update_io.seek(0)
    
    uploaded_update_file = SimpleUploadedFile("clientes_upd.xlsx", excel_update_io.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    res_update = parse_and_import_clients(uploaded_update_file, business, default_prefix='+54', duplicate_mode='update')
    print(f"Resultado duplicados en update: {res_update}")
    assert res_update['updated'] == 1, f"Debió actualizar 1 cliente, actualizó: {res_update['updated']}"
    
    c1.refresh_from_db()
    assert c1.address == 'NUEVA DIRECCION 999', f"Dirección no actualizada: {c1.address}"
    print(f"[OK] Cliente actualizado: {c1.full_name}, Nueva dirección: {c1.address}")
    print("[PASSED] Test 4: Modo 'update' actualizó los datos existentes sin crear duplicados.")

    # 6. Test Importación desde CSV
    print("\n--- Test 5: Importación desde archivo CSV ---")
    csv_content = """Nombres;Apellidos;Celular;Correo;Identificacion;Domicilio;Status;Observacion
Mateo;Gómez;+5491100030003;mateo@test.com;45123456;Recoleta 444;Activo;Cliente de prueba CSV
Sofía;López;+5491100040004;sofia@test.com;46789123;Belgrano 888;Prospecto;Interesada en spa
"""
    uploaded_csv = SimpleUploadedFile("clientes.csv", csv_content.encode('utf-8'), content_type="text/csv")
    res_csv = parse_and_import_clients(uploaded_csv, business, default_prefix='+54', duplicate_mode='skip')
    print(f"Resultado importación CSV: {res_csv}")
    assert res_csv['created'] == 2, f"Se esperaban 2 creados desde CSV, se obtuvieron {res_csv['created']}"
    
    c_csv = Client.objects.get(business=business, phone='+5491100030003')
    assert c_csv.first_name == 'Mateo'
    assert c_csv.last_name == 'Gómez'
    print(f"[OK] Cliente desde CSV: {c_csv.full_name} ({c_csv.phone})")
    print("[PASSED] Test 5: Archivo CSV parseado e importado con éxito.")

    # Limpieza final
    Client.objects.filter(business=business, phone__in=test_phones).delete()
    print("\n=== TODAS LAS PRUEBAS PASARON EXITOSAMENTE (5/5) ===")

if __name__ == '__main__':
    run_tests()
