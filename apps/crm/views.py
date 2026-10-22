import csv
import io
import unicodedata
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.db import transaction
from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from apps.business.models import Business
from apps.crm.models import Client, ClientNote, Supplier
from apps.core.utils import split_phone_number, format_phone_with_prefix

@login_required
def client_list_view(request):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    query = request.GET.get('q', '')
    real_estate_filter = request.GET.get('re_filter', 'all')
    
    is_real_estate = business.is_real_estate and business.is_module_enabled('properties')
    
    clients = Client.objects.filter(business=business).order_by('-created_at') if business else []
    if is_real_estate:
        clients = clients.prefetch_related('owned_properties', 'property_inquiries', 'appointments__related_property')
        
    if query:
        clients = clients.filter(first_name__icontains=query) | clients.filter(last_name__icontains=query) | clients.filter(phone__icontains=query)

    if is_real_estate and real_estate_filter and real_estate_filter != 'all':
        from django.db.models import Q
        seller_q = (
            Q(owned_properties__isnull=False) |
            Q(property_inquiries__inquiry_type__in=['SELLER', 'SELL_TO_BUY']) |
            Q(notes__icontains='[PROPIETARIO')
        )
        buyer_q = (
            Q(property_inquiries__inquiry_type__in=['BUYER', 'SELL_TO_BUY']) |
            Q(appointments__related_property__isnull=False)
        )
        if real_estate_filter == 'sellers':
            clients = clients.filter(seller_q).distinct()
        elif real_estate_filter == 'buyers':
            clients = clients.filter(buyer_q).distinct()
        elif real_estate_filter == 'sell_to_buy':
            clients = clients.filter(
                Q(property_inquiries__inquiry_type='SELL_TO_BUY') |
                Q(notes__icontains='VENDER PARA COMPRAR') |
                (
                    (Q(owned_properties__isnull=False) | Q(property_inquiries__inquiry_type='SELLER') | Q(notes__icontains='[PROPIETARIO')) &
                    (Q(property_inquiries__inquiry_type='BUYER') | Q(appointments__related_property__isnull=False))
                )
            ).distinct()
        elif real_estate_filter == 'visits':
            clients = clients.filter(appointments__related_property__isnull=False).distinct()

    from django.core.paginator import Paginator
    paginator = Paginator(clients, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    return render(request, 'crm/client_list.html', {
        'business': business,
        'clients': page_obj,
        'page_obj': page_obj,
        'query': query,
        'is_real_estate': is_real_estate,
        'real_estate_filter': real_estate_filter,
    })

@login_required
def client_create_view(request):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    
    default_prefix, _ = split_phone_number(business.phone if business else '')

    if request.method == 'POST':
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        phone_prefix = request.POST.get('phone_prefix', default_prefix).strip()
        phone_number = request.POST.get('phone_number', '').strip()
        
        phone = format_phone_with_prefix(phone_prefix, phone_number, default_prefix=default_prefix) or "Sin teléfono"

        email = request.POST.get('email', '')
        tax_id = request.POST.get('tax_id', '')
        address = request.POST.get('address', '')
        status = request.POST.get('status', 'ACTIVE')
        notes = request.POST.get('notes', '')
        
        Client.objects.create(
            business=business,
            first_name=first_name,
            last_name=last_name,
            phone=phone,
            email=email,
            tax_id=tax_id,
            address=address,
            status=status,
            notes=notes
        )
        return redirect('client_list')
        
    return render(request, 'crm/client_form.html', {
        'business': business,
        'client': None,
        'phone_prefix': default_prefix,
        'phone_number': '',
        'title': 'Crear Nuevo Cliente'
    })

@login_required
def client_edit_view(request, client_id):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    client = get_object_or_404(Client, id=client_id, business=business)
    default_prefix, _ = split_phone_number(business.phone if business else '')

    if request.method == 'POST':
        client.first_name = request.POST.get('first_name', '').strip()
        client.last_name = request.POST.get('last_name', '').strip()
        
        phone_prefix = request.POST.get('phone_prefix', default_prefix).strip()
        phone_number = request.POST.get('phone_number', '').strip()
        client.phone = format_phone_with_prefix(phone_prefix, phone_number, default_prefix=default_prefix) or "Sin teléfono"

        client.email = request.POST.get('email', '')
        client.tax_id = request.POST.get('tax_id', '')
        client.address = request.POST.get('address', '')
        client.status = request.POST.get('status', 'ACTIVE')
        client.notes = request.POST.get('notes', '')
        client.save()
        return redirect('client_detail', client_id=client.id)

    phone_prefix, phone_number = split_phone_number(client.phone)
        
    return render(request, 'crm/client_form.html', {
        'business': business,
        'client': client,
        'phone_prefix': phone_prefix,
        'phone_number': phone_number,
        'title': f'Editar Cliente: {client.full_name}'
    })

from apps.crm.models import Client, ClientNote, Supplier, ServiceHistoryNote
from django.contrib import messages

def parse_legacy_seller_notes(notes):
    if not notes or '[PROPIETARIO' not in notes:
        return []
    results = []
    for line in notes.split('\n'):
        line_clean = line.strip()
        if '[PROPIETARIO' in line_clean:
            data = {
                'operation_type': 'Venta',
                'property_type': 'Inmueble',
                'address_zone': '',
                'rooms': '',
                'estimated_price': '',
                'comments': '',
                'raw_line': line_clean
            }
            idx = line_clean.find(']:')
            content_part = line_clean[idx+2:].strip() if idx != -1 else line_clean
            parts = content_part.split(' · ')
            for p in parts:
                if ':' in p:
                    k, v = p.split(':', 1)
                    k, v = k.strip().lower(), v.strip()
                    if 'operación' in k:
                        data['operation_type'] = v
                    elif 'tipo' in k:
                        data['property_type'] = v
                    elif 'ubicación' in k:
                        data['address_zone'] = v
                    elif 'ambientes' in k:
                        data['rooms'] = v
                    elif 'precio' in k:
                        data['estimated_price'] = v
                    elif 'comentarios' in k:
                        data['comments'] = v
            results.append(data)
    return results

@login_required
def client_detail_view(request, client_id):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    client = get_object_or_404(Client, id=client_id, business=business)
    
    appointments = client.appointments.select_related('staff', 'service', 'related_property').order_by('-date', '-start_time')
    sales = client.sales.select_related('staff').order_by('-created_at')
    invoices = client.invoices.order_by('-issue_date')
    notes = client.client_notes.all().order_by('-created_at')
    service_notes = client.service_notes.select_related('appointment').order_by('-created_at')
    
    allow_editing_history = getattr(business, 'allow_editing_client_history', True)
    
    # Real estate intelligence
    is_real_estate = business.is_real_estate and business.is_module_enabled('properties')
    
    seller_inquiries = []
    buyer_inquiries = []
    sell_to_buy_inquiries = []
    property_visits = []
    owned_properties = []
    available_properties = []
    staff_members = []
    legacy_seller_requests = []
    
    if is_real_estate:
        from apps.properties.models import Property, PropertyInquiry
        owned_properties = client.owned_properties.select_related('agent').order_by('-created_at')
        seller_inquiries = client.property_inquiries.filter(inquiry_type='SELLER').select_related('agent').order_by('-created_at')
        buyer_inquiries = client.property_inquiries.filter(inquiry_type='BUYER').select_related('property', 'agent').order_by('-created_at')
        sell_to_buy_inquiries = client.property_inquiries.filter(inquiry_type='SELL_TO_BUY').select_related('property', 'agent').order_by('-created_at')
        property_visits = [app for app in appointments if app.related_property]
        legacy_seller_requests = parse_legacy_seller_notes(client.notes)
        available_properties = Property.objects.filter(business=business, status='AVAILABLE').order_by('title')
        staff_members = business.staff_members.filter(is_active=True).order_by('first_name', 'last_name')
    
    return render(request, 'crm/client_detail.html', {
        'business': business,
        'client': client,
        'appointments': appointments,
        'sales': sales,
        'invoices': invoices,
        'notes': notes,
        'service_notes': service_notes,
        'allow_editing_history': allow_editing_history,
        'is_real_estate': is_real_estate,
        'seller_inquiries': seller_inquiries,
        'buyer_inquiries': buyer_inquiries,
        'sell_to_buy_inquiries': sell_to_buy_inquiries,
        'property_visits': property_visits,
        'owned_properties': owned_properties,
        'available_properties': available_properties,
        'staff_members': staff_members,
        'legacy_seller_requests': legacy_seller_requests,
    })

@login_required
def add_client_property_inquiry_view(request, client_id):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    client = get_object_or_404(Client, id=client_id, business=business)
    
    if request.method == 'POST':
        from apps.properties.models import Property, PropertyInquiry
        inquiry_type = request.POST.get('inquiry_type', 'BUYER').strip()
        operation_type = request.POST.get('operation_type', 'Venta').strip()
        property_type = request.POST.get('property_type', 'Departamento').strip()
        address_zone = request.POST.get('address_zone', '').strip()
        rooms = request.POST.get('rooms', '').strip()
        estimated_price = request.POST.get('estimated_price', '').strip()
        property_id = request.POST.get('property_id')
        notes = request.POST.get('notes', '').strip()
        agent_id = request.POST.get('agent_id')
        status = request.POST.get('status', 'PENDING')

        if inquiry_type == 'SELL_TO_BUY':
            operation_type = 'Vender para Comprar'
            target_property_type = request.POST.get('target_property_type', '').strip()
            target_zone = request.POST.get('target_zone', '').strip()
            target_rooms = request.POST.get('target_rooms', '').strip()
            target_budget = request.POST.get('target_budget', '').strip()
            target_parts = []
            if target_property_type:
                target_parts.append(f"Busca comprar: {target_property_type}")
            if target_zone:
                target_parts.append(f"Zona buscada: {target_zone}")
            if target_rooms:
                target_parts.append(f"Ambientes buscados: {target_rooms}")
            if target_budget:
                target_parts.append(f"Presupuesto / Diferencia: {target_budget}")
            if target_parts:
                target_summary = " | ".join(target_parts)
                notes = f"[PROPIEDAD QUE BUSCA COMPRAR] {target_summary}\n{notes}".strip() if notes else f"[PROPIEDAD QUE BUSCA COMPRAR] {target_summary}"
        
        selected_property = None
        if property_id:
            selected_property = Property.objects.filter(id=property_id, business=business).first()
            
        agent = None
        if agent_id:
            agent = business.staff_members.filter(id=agent_id).first()
            
        PropertyInquiry.objects.create(
            business=business,
            client=client,
            property=selected_property,
            inquiry_type=inquiry_type,
            operation_type=operation_type,
            property_type=property_type,
            address_zone=address_zone,
            rooms=rooms or None,
            estimated_price=estimated_price or None,
            notes=notes,
            status=status,
            agent=agent
        )
        
        if inquiry_type == 'SELL_TO_BUY':
            label = "Operación simultánea (Vender para Comprar)"
        elif inquiry_type == 'BUYER':
            label = "Interés de compra / alquiler"
        else:
            label = "Solicitud de propietario"
        messages.success(request, f"{label} registrada exitosamente para {client.full_name}.")
        
    return redirect('client_detail', client_id=client.id)

@login_required
def update_property_inquiry_status_view(request, inquiry_id):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    from apps.properties.models import PropertyInquiry
    inquiry = get_object_or_404(PropertyInquiry, id=inquiry_id, business=business)
    
    if request.method == 'POST':
        new_status = request.POST.get('status')
        agent_id = request.POST.get('agent_id')
        if new_status and new_status in dict(PropertyInquiry.STATUS_CHOICES):
            inquiry.status = new_status
        if agent_id is not None:
            if agent_id == "":
                inquiry.agent = None
            else:
                inquiry.agent = business.staff_members.filter(id=agent_id).first()
        inquiry.save()
        messages.success(request, "Estado de la solicitud inmobiliaria actualizado correctamente.")
        
    return redirect('client_detail', client_id=inquiry.client.id)

@login_required
def delete_property_inquiry_view(request, inquiry_id):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    from apps.properties.models import PropertyInquiry
    inquiry = get_object_or_404(PropertyInquiry, id=inquiry_id, business=business)
    client_id = inquiry.client.id
    
    if request.method == 'POST':
        inquiry.delete()
        messages.success(request, "Solicitud inmobiliaria eliminada.")
        
    return redirect('client_detail', client_id=client_id)

@login_required
def add_client_service_note_view(request, client_id):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    client = get_object_or_404(Client, id=client_id, business=business)
    
    if not getattr(business, 'allow_editing_client_history', True):
        messages.warning(request, "La edición del historial de servicios está bloqueada en la configuración de la empresa.")
        return redirect('client_detail', client_id=client.id)
        
    if request.method == 'POST':
        service_name = request.POST.get('service_name', '').strip()
        content = request.POST.get('content', '').strip()
        appointment_id = request.POST.get('appointment_id')
        author = request.user.get_full_name() or request.user.username if request.user.is_authenticated else "Staff"
        
        if service_name and content:
            ServiceHistoryNote.objects.create(
                client=client,
                service_name=service_name,
                content=content,
                author=author,
                appointment_id=appointment_id if appointment_id else None
            )
            messages.success(request, "Anotación de servicio agregada a la ficha del cliente.")
        else:
            messages.error(request, "Por favor ingresa el servicio y el detalle de la anotación.")
            
    return redirect('client_detail', client_id=client.id)

@login_required
def delete_client_service_note_view(request, note_id):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    note = get_object_or_404(ServiceHistoryNote, id=note_id, client__business=business)
    client_id = note.client.id
    
    if not getattr(business, 'allow_editing_client_history', True):
        messages.warning(request, "No se pueden eliminar anotaciones porque el historial está en modo Solo Lectura.")
        return redirect('client_detail', client_id=client_id)
        
    note.delete()
    messages.success(request, "Anotación de servicio eliminada.")
    return redirect('client_detail', client_id=client_id)

@login_required
def update_client_general_notes_view(request, client_id):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    client = get_object_or_404(Client, id=client_id, business=business)
    if request.method == 'POST':
        notes = request.POST.get('notes', '').strip()
        client.notes = notes
        client.save(update_fields=['notes'])
        messages.success(request, "Notas y observaciones generales actualizadas con éxito.")
    return redirect('client_detail', client_id=client.id)

@login_required
def supplier_list_view(request):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    suppliers = Supplier.objects.filter(business=business).order_by('-created_at') if business else []

    from django.core.paginator import Paginator
    paginator = Paginator(suppliers, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    return render(request, 'crm/supplier_list.html', {
        'business': business,
        'suppliers': page_obj,
        'page_obj': page_obj,
    })

@login_required
def supplier_create_view(request):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    
    default_prefix, _ = split_phone_number(business.phone if business else '')

    if request.method == 'POST':
        company_name = request.POST.get('company_name')
        contact_name = request.POST.get('contact_name', '')
        tax_id = request.POST.get('tax_id', '')
        email = request.POST.get('email', '')
        phone_prefix = request.POST.get('phone_prefix', default_prefix).strip()
        phone_number = request.POST.get('phone_number', '').strip() or request.POST.get('phone', '').strip()
        phone = format_phone_with_prefix(phone_prefix, phone_number, default_prefix=default_prefix)
        address = request.POST.get('address', '')
        category = request.POST.get('category', 'Insumos Generales')
        
        Supplier.objects.create(
            business=business,
            company_name=company_name,
            contact_name=contact_name,
            tax_id=tax_id,
            email=email,
            phone=phone,
            address=address,
            category=category
        )
        return redirect('supplier_list')
        
    return render(request, 'crm/supplier_form.html', {
        'business': business,
        'supplier': None,
        'phone_prefix': default_prefix,
        'phone_number': '',
        'title': 'Crear Nuevo Proveedor'
    })

@login_required
def supplier_edit_view(request, supplier_id):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    supplier = get_object_or_404(Supplier, id=supplier_id, business=business)
    default_prefix, _ = split_phone_number(business.phone if business else '')

    if request.method == 'POST':
        supplier.company_name = request.POST.get('company_name')
        supplier.contact_name = request.POST.get('contact_name', '')
        supplier.tax_id = request.POST.get('tax_id', '')
        supplier.email = request.POST.get('email', '')
        phone_prefix = request.POST.get('phone_prefix', default_prefix).strip()
        phone_number = request.POST.get('phone_number', '').strip() or request.POST.get('phone', '').strip()
        supplier.phone = format_phone_with_prefix(phone_prefix, phone_number, default_prefix=default_prefix)
        supplier.address = request.POST.get('address', '')
        supplier.category = request.POST.get('category', 'Insumos Generales')
        supplier.save()
        return redirect('supplier_list')
        
    supp_prefix, supp_num = split_phone_number(supplier.phone, default_prefix=default_prefix)

    return render(request, 'crm/supplier_form.html', {
        'business': business,
        'supplier': supplier,
        'phone_prefix': supp_prefix,
        'phone_number': supp_num,
        'title': 'Editar Proveedor'
    })

@login_required
def export_clients_excel(request):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    clients = Client.objects.filter(business=business) if business else []
    
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="clientes_reporte.csv"'
    
    response.write('\ufeff')
    writer = csv.writer(response)
    writer.writerow(['Nombre', 'Apellido', 'Teléfono', 'Email', 'Documento/Tax ID', 'Dirección', 'Estado', 'Fecha Registro'])
    
    for c in clients:
        writer.writerow([
            c.first_name,
            c.last_name,
            c.phone or '',
            c.email or '',
            c.tax_id or '',
            c.address or '',
            c.status,
            c.created_at.strftime('%Y-%m-%d %H:%M') if hasattr(c, 'created_at') and c.created_at else ''
        ])
        
    return response

@login_required
def export_clients_pdf(request):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')
    clients = Client.objects.filter(business=business) if business else []
    
    rows = []
    for c in clients:
        rows.append({
            'col1': c.full_name,
            'col2': c.phone or 'Sin teléfono',
            'col3': c.email or '-',
            'col4': c.address or '-',
            'col5': c.status
        })
        
    return render(request, 'analytics/pdf_report.html', {
        'title': 'Reporte de Clientes',
        'subtitle': f'Base de Datos de Clientes - {business.name if business else ""}',
        'business': business,
        'headers': ['Cliente', 'Teléfono', 'Email', 'Dirección', 'Estado'],
        'rows': rows,
        'summary': f'Total Clientes: {len(rows)}'
    })


def _normalize_header(header_text):
    if not header_text:
        return ""
    text = str(header_text).strip()
    return "".join(
        c for c in unicodedata.normalize('NFD', text)
        if unicodedata.category(c) != 'Mn'
    ).lower().replace('_', ' ').replace('-', ' ').strip()


def _clean_phone_number(raw_val, default_prefix='+54'):
    if raw_val is None:
        return 'Sin teléfono'
    if isinstance(raw_val, float):
        raw_val = int(raw_val)
    val_str = str(raw_val).strip()
    if not val_str or val_str.lower() in ['nan', 'none', 'null', '-', 'sin telefono', 'sin teléfono']:
        return 'Sin teléfono'
        
    if val_str.startswith('+'):
        clean = '+' + ''.join(filter(str.isdigit, val_str))
        return clean if len(clean) > 1 else 'Sin teléfono'
        
    digits = ''.join(filter(str.isdigit, val_str))
    if not digits:
        return 'Sin teléfono'
        
    if digits.startswith('00'):
        return '+' + digits[2:]
        
    prefix = default_prefix.strip()
    if not prefix.startswith('+'):
        prefix = f"+{prefix}"
        
    prefix_digits = prefix.replace('+', '')
    if prefix_digits and digits.startswith(prefix_digits):
        return f"+{digits}"
        
    return f"{prefix}{digits}"


def _clean_status(raw_val):
    if not raw_val:
        return 'ACTIVE'
    val = _normalize_header(raw_val)
    if 'vip' in val or 'frecuente' in val:
        return 'VIP'
    elif 'inactiv' in val:
        return 'INACTIVE'
    elif 'prospect' in val or 'potencial' in val:
        return 'PROSPECT'
    return 'ACTIVE'


def _clean_date(raw_val):
    if not raw_val:
        return None
    if isinstance(raw_val, (datetime.date, datetime.datetime)):
        return raw_val.date() if isinstance(raw_val, datetime.datetime) else raw_val
    val_str = str(raw_val).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
        try:
            return datetime.datetime.strptime(val_str, fmt).date()
        except ValueError:
            pass
    return None


def _extract_rows_from_file(uploaded_file):
    filename = uploaded_file.name.lower()
    rows_data = []
    
    if filename.endswith(('.xlsx', '.xlsm', '.xltx')):
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        sheet = wb.active
        for row in sheet.iter_rows(values_only=True):
            if row and any(cell is not None and str(cell).strip() != '' for cell in row):
                rows_data.append([cell if cell is not None else '' for cell in row])
    elif filename.endswith(('.csv', '.txt')):
        content = uploaded_file.read()
        uploaded_file.seek(0)
        text = None
        for encoding in ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252']:
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = content.decode('utf-8', errors='ignore')
            
        first_line = text.splitlines()[0] if text.splitlines() else ''
        delimiter = ';' if ';' in first_line and ',' not in first_line else ','
        if '\t' in first_line and ',' not in first_line and ';' not in first_line:
            delimiter = '\t'
            
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        for row in reader:
            if row and any(str(cell).strip() != '' for cell in row):
                rows_data.append([str(cell).strip() for cell in row])
    else:
        raise ValueError("Formato de archivo no soportado. Por favor sube un archivo Excel (.xlsx) o CSV (.csv).")
        
    return rows_data


def parse_and_import_clients(uploaded_file, business, default_prefix='+54', duplicate_mode='skip'):
    rows = _extract_rows_from_file(uploaded_file)
    if not rows:
        return {'created': 0, 'updated': 0, 'skipped': 0, 'errors': ['El archivo está vacío o no contiene filas con datos.']}
        
    column_synonyms = {
        'first_name': ['nombre', 'nombres', 'firstname', 'first name', 'name', 'nombre cliente', 'nombre de cliente', 'cliente', 'contacto'],
        'last_name': ['apellido', 'apellidos', 'lastname', 'last name', 'surname'],
        'phone': ['telefono', 'telefonos', 'whatsapp', 'celular', 'movil', 'phone', 'telephone', 'mobile', 'tel', 'contacto telefonico'],
        'email': ['email', 'correo', 'correo electronico', 'mail', 'e mail'],
        'tax_id': ['dni', 'rut', 'cedula', 'documento', 'identificacion', 'tax id', 'tax_id', 'rfc', 'cuit', 'cuil', 'nif', 'pasaporte', 'numero de identidad'],
        'address': ['direccion', 'domicilio', 'address', 'ubicacion', 'residencia'],
        'status': ['estado', 'status', 'categoria', 'tipo'],
        'notes': ['notas', 'nota', 'observaciones', 'observacion', 'comentarios', 'comentario', 'notes', 'preferencias', 'detalle'],
        'birth_date': ['fecha de nacimiento', 'cumpleanos', 'nacimiento', 'birth date', 'cumple']
    }
    
    header_idx = -1
    col_mapping = {}
    
    for r_idx, row in enumerate(rows[:3]):
        temp_map = {}
        for c_idx, cell in enumerate(row):
            norm_cell = _normalize_header(cell)
            for field, synonyms in column_synonyms.items():
                if field not in temp_map:
                    for syn in synonyms:
                        if norm_cell == syn or (len(norm_cell) > 2 and norm_cell.startswith(syn)):
                            temp_map[field] = c_idx
                            break
        if 'first_name' in temp_map or ('phone' in temp_map and len(temp_map) >= 2):
            header_idx = r_idx
            col_mapping = temp_map
            break
            
    data_rows = []
    if header_idx != -1:
        data_rows = rows[header_idx + 1:]
    else:
        col_mapping = {
            'first_name': 0,
            'last_name': 1 if len(rows[0]) > 1 else None,
            'phone': 2 if len(rows[0]) > 2 else None,
            'email': 3 if len(rows[0]) > 3 else None,
            'tax_id': 4 if len(rows[0]) > 4 else None,
            'address': 5 if len(rows[0]) > 5 else None,
            'status': 6 if len(rows[0]) > 6 else None,
            'notes': 7 if len(rows[0]) > 7 else None,
        }
        data_rows = rows[1:] if len(rows) > 1 else rows

    created_count = 0
    updated_count = 0
    skipped_count = 0
    errors = []

    with transaction.atomic():
        for row_num, row in enumerate(data_rows, start=(header_idx + 2 if header_idx != -1 else 2)):
            def get_val(field):
                idx = col_mapping.get(field)
                if idx is not None and idx < len(row):
                    val = row[idx]
                    return str(val).strip() if val is not None else ''
                return ''

            raw_first_name = get_val('first_name')
            raw_last_name = get_val('last_name')
            raw_phone = get_val('phone')
            raw_email = get_val('email')
            raw_tax_id = get_val('tax_id')
            raw_address = get_val('address')
            raw_status = get_val('status')
            raw_notes = get_val('notes')
            raw_birth_date = get_val('birth_date')

            if not raw_first_name and not raw_last_name:
                continue

            first_name = raw_first_name
            last_name = raw_last_name
            if first_name and not last_name:
                parts = first_name.split(maxsplit=1)
                if len(parts) == 2:
                    first_name, last_name = parts[0], parts[1]

            phone = _clean_phone_number(raw_phone, default_prefix)
            email = raw_email if '@' in raw_email else ('' if not raw_email else None)
            tax_id = raw_tax_id or ''
            address = raw_address or ''
            status = _clean_status(raw_status)
            notes = raw_notes or ''
            birth_date = _clean_date(raw_birth_date)

            existing_client = None
            if duplicate_mode in ['skip', 'update']:
                query = Client.objects.filter(business=business)
                if phone and phone != 'Sin teléfono':
                    existing_client = query.filter(phone=phone).first()
                if not existing_client and email:
                    existing_client = query.filter(email__iexact=email).first()
                if not existing_client and tax_id:
                    existing_client = query.filter(tax_id=tax_id).first()

            if existing_client:
                if duplicate_mode == 'skip':
                    skipped_count += 1
                    continue
                elif duplicate_mode == 'update':
                    if first_name:
                        existing_client.first_name = first_name
                    if last_name:
                        existing_client.last_name = last_name
                    if phone and phone != 'Sin teléfono':
                        existing_client.phone = phone
                    if email:
                        existing_client.email = email
                    if tax_id:
                        existing_client.tax_id = tax_id
                    if address:
                        existing_client.address = address
                    if raw_status:
                        existing_client.status = status
                    if notes:
                        existing_client.notes = f"{existing_client.notes}\n{notes}".strip() if existing_client.notes else notes
                    if birth_date:
                        existing_client.birth_date = birth_date
                    existing_client.save()
                    updated_count += 1
                    continue

            Client.objects.create(
                business=business,
                first_name=first_name,
                last_name=last_name or '',
                phone=phone,
                email=email or None,
                tax_id=tax_id or None,
                address=address or None,
                status=status,
                notes=notes or None,
                birth_date=birth_date
            )
            created_count += 1

    return {
        'created': created_count,
        'updated': updated_count,
        'skipped': skipped_count,
        'errors': errors
    }


@login_required
def client_import_view(request):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')

    if request.method != 'POST':
        return redirect('client_list')

    uploaded_file = request.FILES.get('file')
    if not uploaded_file:
        messages.error(request, "Por favor selecciona un archivo Excel (.xlsx) o CSV para importar.")
        return redirect('client_list')

    default_prefix = request.POST.get('default_prefix', '+54').strip()
    duplicate_mode = request.POST.get('duplicate_mode', 'skip').strip()

    try:
        result = parse_and_import_clients(
            uploaded_file=uploaded_file,
            business=business,
            default_prefix=default_prefix,
            duplicate_mode=duplicate_mode
        )

        created = result['created']
        updated = result['updated']
        skipped = result['skipped']
        errors = result.get('errors', [])

        if created > 0 or updated > 0:
            summary_parts = []
            if created > 0:
                summary_parts.append(f"{created} nuevo{'s' if created > 1 else ''}")
            if updated > 0:
                summary_parts.append(f"{updated} actualizado{'s' if updated > 1 else ''}")
            if skipped > 0:
                summary_parts.append(f"{skipped} duplicado{'s' if skipped > 1 else ''} omitido{'s' if skipped > 1 else ''}")

            messages.success(request, f"¡Importación completada! Se procesaron {', '.join(summary_parts)}.")
        elif skipped > 0:
            messages.info(request, f"Proceso finalizado: se omitieron {skipped} clientes porque ya se encontraban registrados en la base de datos.")
        else:
            messages.warning(request, "No se crearon clientes. Asegúrate de que el archivo tenga al menos una columna de 'Nombre' con datos válidos.")

        for err in errors[:3]:
            messages.warning(request, err)

    except Exception as e:
        messages.error(request, f"Error al procesar el archivo de importación: {str(e)}")

    return redirect('client_list')


@login_required
def client_template_download_view(request):
    business = getattr(request, 'current_business', None) or request.user.business
    if not business:
        return redirect('onboarding')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clientes"

    headers = [
        "Nombre *",
        "Apellido",
        "Teléfono / WhatsApp",
        "Email",
        "DNI / Documento",
        "Dirección",
        "Estado (Activo/VIP/Prospecto/Inactivo)",
        "Notas"
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="881337", end_color="881337", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border_thin = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)

    ws.row_dimensions[1].height = 28
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = cell_border

    sample_data = [
        ["María", "González", "+54 9 11 2345-6789", "maria.gonzalez@ejemplo.com", "34567890", "Av. Corrientes 1234, CABA", "Activo", "Prefiere turnos por la tarde"],
        ["Carlos", "Rodríguez", "+54 9 11 9876-5432", "carlos.rod@ejemplo.com", "28901234", "Calle San Martín 456", "VIP", "Cliente frecuente"],
        ["Lucía", "Fernández", "+54 9 11 4455-6677", "lucia.f@ejemplo.com", "39123456", "Belgrano 789", "Prospecto", "Consulta por promociones y nuevos servicios"],
    ]

    for row_idx, row_values in enumerate(sample_data, start=2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = cell_border
            cell.alignment = Alignment(vertical="center")

    column_widths = [18, 18, 24, 30, 20, 32, 24, 35]
    for col_idx, width in enumerate(column_widths, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    ws_instructions = wb.create_sheet(title="Instrucciones")
    ws_instructions.cell(row=1, column=1, value="Guía de Importación de Clientes").font = Font(size=14, bold=True, color="881337")
    instructions = [
        "1. El único campo estrictamente obligatorio es 'Nombre *'.",
        "2. Si colocas el nombre y apellido juntos en la columna 'Nombre' y dejas 'Apellido' vacío, el sistema los dividirá automáticamente.",
        "3. El teléfono puede incluir o no el código de país. En el modal de importación podrás seleccionar el prefijo predeterminado (+54, +56, +57, etc.).",
        "4. Estados reconocidos: 'Activo', 'VIP', 'Prospecto' o 'Inactivo'. Si lo dejas en blanco, se asignará 'Activo' por defecto.",
        "5. Puedes omitir o eliminar las filas de ejemplo antes de subir tu archivo con los clientes reales.",
        "6. Formatos soportados en la plataforma: Excel (.xlsx) y archivos de texto (.csv)."
    ]
    for i, inst in enumerate(instructions, start=3):
        ws_instructions.cell(row=i, column=1, value=inst).font = Font(size=11)
    ws_instructions.column_dimensions['A'].width = 85

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Plantilla_Importar_Clientes.xlsx"'
    return response
